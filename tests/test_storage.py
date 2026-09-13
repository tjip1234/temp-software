"""Recording, reading back, and exporting."""

from __future__ import annotations

import asyncio

import numpy as np
import pytest

from tjiptemp.device.device import Device
from tjiptemp.protocol.channels import Ch
from tjiptemp.protocol.messages import (
    BLOCK_FLAG_DECIMATED,
    SampleBlock,
    decode_sample_block,
    encode_sample_block,
)
from tjiptemp.simulator.board import SimulatedBoard
from tjiptemp.simulator.server import LoopbackTransport, SimulatorRuntime
from tjiptemp.storage.db import Database, pack_block
from tjiptemp.storage.export import ExportOptions, summary_table, to_csv, to_excel, to_json
from tjiptemp.storage.recorder import Recorder


@pytest.fixture
def db(tmp_path):
    database = Database(tmp_path / "test.tjip")
    yield database
    database.close()


@pytest.fixture
async def runtime():
    rt = SimulatorRuntime(SimulatedBoard(serial="TJIP-STORE0001", rate_hz=50.0, seed=3))
    await rt.start()
    yield rt
    await rt.stop()


# ------------------------------------------------------------------- database

def test_schema_and_stats(db):
    stats = db.stats()
    assert stats["sessions"] == 0
    assert stats["rows"] == 0


def test_blocks_roundtrip_through_sqlite(db):
    db.upsert_device("TJIP-X", {"model": "test"})
    session = db.create_session("TJIP-X", [0, 1, 2], rate_hz=10.0)

    values = np.arange(30, dtype=np.float32).reshape(10, 3)
    db.append_blocks(session, [pack_block(0, values, t0_utc=1000.0, dt_s=0.1)])

    times, read_back, channels = db.read_session(session)
    assert channels == [0, 1, 2]
    np.testing.assert_allclose(read_back, values)
    np.testing.assert_allclose(times, 1000.0 + np.arange(10) * 0.1)
    assert db.count_rows(session) == 10


def test_backfilled_blocks_are_sorted_on_read(db):
    """Backfill arrives after newer data; the export must still be in time order."""
    db.upsert_device("TJIP-X", {})
    session = db.create_session("TJIP-X", [0])

    late = np.full((5, 1), 2.0, dtype=np.float32)
    early = np.full((5, 1), 1.0, dtype=np.float32)
    db.append_blocks(session, [pack_block(10, late, t0_utc=2000.0, dt_s=1.0)])
    db.append_blocks(session, [pack_block(0, early, t0_utc=1000.0, dt_s=1.0)])

    times, values, _ = db.read_session(session)
    assert np.all(np.diff(times) > 0)
    assert values[0, 0] == 1.0 and values[-1, 0] == 2.0


def test_time_range_query(db):
    db.upsert_device("TJIP-X", {})
    session = db.create_session("TJIP-X", [0])
    values = np.arange(100, dtype=np.float32).reshape(100, 1)
    db.append_blocks(session, [pack_block(0, values, t0_utc=0.0, dt_s=1.0)])

    times, sliced, _ = db.read_session(session, t_from=20.0, t_to=29.0)
    assert len(times) == 10
    assert sliced[0, 0] == 20.0 and sliced[-1, 0] == 29.0


def test_calibration_history_is_kept(db):
    db.upsert_device("TJIP-X", {})
    db.record_calibration("TJIP-X", {"rev": 1, "by": "raaf", "channels": {}})
    db.record_calibration("TJIP-X", {"rev": 2, "by": "raaf", "channels": {}})
    history = db.calibration_history("TJIP-X")
    assert [h["rev"] for h in history] == [2, 1]


def test_newer_schema_is_refused_rather_than_guessed(tmp_path):
    path = tmp_path / "future.tjip"
    database = Database(path)
    with database.connect() as conn:
        conn.execute("UPDATE meta SET value = '99' WHERE key = 'schema_version'")
    database.close()

    with pytest.raises(RuntimeError, match="newer version"):
        Database(path)


# ------------------------------------------------------------------- recorder

async def test_record_a_live_session_and_read_it_back(db, runtime):
    device = Device()
    await device.add_transport(LoopbackTransport(runtime))
    recorder = Recorder(db)
    try:
        await device.start_streaming(50.0)
        await asyncio.sleep(0.3)
        recording = recorder.start(device, name="soak test")
        await asyncio.sleep(1.5)
        recorder.mark(device.serial, "lid closed")
        await recorder.stop(device.serial)

        info = db.get_session(recording.session_id)
        assert info is not None
        assert info.state == "complete"
        assert info.rows > 20
        assert info.cal, "the calibration in force must be snapshotted"

        times, values, channels = db.read_session(recording.session_id)
        assert len(times) == info.rows
        assert len(channels) == 15
        assert np.all(np.diff(times) > 0)

        pt1000 = values[:, channels.index(int(Ch.PT1000))]
        assert np.all(np.isfinite(pt1000))
        assert 10.0 < float(np.mean(pt1000)) < 60.0

        assert len(db.list_markers(recording.session_id)) == 1
    finally:
        await recorder.stop_all()
        await device.close()


async def test_recording_twice_is_refused(db, runtime):
    device = Device()
    await device.add_transport(LoopbackTransport(runtime))
    recorder = Recorder(db)
    try:
        await device.start_streaming(20.0)
        await asyncio.sleep(0.2)
        recorder.start(device)
        with pytest.raises(RuntimeError, match="already recording"):
            recorder.start(device)
    finally:
        await recorder.stop_all()
        await device.close()


# ------------------------------------------------------------------ decimation

def test_decimated_block_roundtrips_with_its_step():
    block = SampleBlock(
        first_seq=100, t0_us=0, dt_us=50_000, channel_ids=(0, 1),
        data=np.zeros((4, 2), dtype=np.float32), seq_step=5,
    )
    decoded = decode_sample_block(encode_sample_block(block))
    assert decoded.seq_step == 5
    assert decoded.is_decimated
    assert decoded.flags & BLOCK_FLAG_DECIMATED
    np.testing.assert_array_equal(decoded.sequences(), [100, 105, 110, 115])
    assert decoded.last_seq == 115


async def test_decimated_blocks_display_but_do_not_manufacture_gaps(runtime):
    """A BLE-style thinned stream must not send the host chasing phantom gaps."""
    device = Device()
    # A 5 Hz subscription against a 50 Hz board: one row in ten.
    await device.add_transport(
        LoopbackTransport(runtime, decimate=True), start_stream=False
    )
    try:
        await device.start_streaming(5.0)
        await asyncio.sleep(1.5)

        assert device.live.size > 3, "decimated rows must still reach the live view"
        assert device.aggregator.stats.rows_accepted == 0
        assert device.aggregator.missing_rows == 0, "no phantom gaps"
        assert device.aggregator.open_gaps == []
    finally:
        await device.close()


# -------------------------------------------------------------------- export

@pytest.fixture
def populated(db):
    db.upsert_device("TJIP-EXPORT", {"model": "tjiptemp-s3", "fw_ver": "0.1.0"})
    session = db.create_session(
        "TJIP-EXPORT", [int(Ch.PT1000), int(Ch.TYPEK), int(Ch.V_BAT)],
        name="export test", rate_hz=10.0,
        cal={"rev": 3, "updated_utc": "2026-08-01T10:00:00Z", "reference": "Fluke 1524"},
    )
    rng = np.random.default_rng(1)
    values = np.column_stack([
        21.0 + rng.normal(0, 0.01, 200),
        22.5 + rng.normal(0, 0.05, 200),
        3.9 + rng.normal(0, 0.001, 200),
    ]).astype(np.float32)
    values[50, 1] = np.nan  # a faulted reading, which must survive as an empty cell
    db.append_blocks(session, [pack_block(0, values, t0_utc=1_800_000_000.0, dt_s=0.1)])
    db.close_session(session, timebase={"drift_ppm": 12.4, "uncertainty_us": 180,
                                        "n_points": 32})
    return session


def test_csv_export_carries_provenance(db, populated, tmp_path):
    path = to_csv(db, populated, tmp_path / "out.csv")
    text = path.read_text(encoding="utf-8")

    assert "# Device serial: TJIP-EXPORT" in text
    assert "# Calibration revision: 3" in text
    assert "# Clock drift: +12.40 ppm" in text
    assert "PT1000 [°C]" in text
    assert "Type K [°C]" in text
    # The faulted sample must be an empty cell, never an interpolated value.
    body = [line for line in text.splitlines() if not line.startswith("#")]
    assert any(",," in line for line in body)


def test_export_names_channels_as_the_recording_board_did(db, tmp_path):
    """The DIN-6's channel 8 is its power area; the built-in table calls id 8 the
    charger thermistor of the earlier board, and exports used to say so."""
    db.upsert_device("TJIP-DIN6", {"model": "tjiptemp-din6", "channels": [
        {"id": 0, "key": "pt1000", "name": "PT1000", "unit": "degC", "kind": "rtd"},
        {"id": 8, "key": "ntc_brd_pwr", "name": "Board: power area", "unit": "degC", "kind": "ntc"},
    ]})
    session = db.create_session("TJIP-DIN6", [0, 8], name="din6", rate_hz=20.0)
    values = np.column_stack([np.full(20, 25.6), np.full(20, 31.4)]).astype(np.float32)
    db.append_blocks(session, [pack_block(0, values, t0_utc=1_800_000_000.0, dt_s=0.05)])
    db.close_session(session)

    text = to_csv(db, session, tmp_path / "din6.csv").read_text(encoding="utf-8")
    assert "Board: power area [°C]" in text
    assert "charger" not in text


def test_a_board_that_declared_no_channels_falls_back_to_the_built_in_names(db, tmp_path):
    db.upsert_device("TJIP-OLDFW", {"model": "tjiptemp-s3", "fw_ver": "0.0.9"})
    session = db.create_session("TJIP-OLDFW", [0], name="orphan", rate_hz=10.0)
    db.append_blocks(session, [pack_block(0, np.full((5, 1), 20.0, np.float32),
                                          t0_utc=1_800_000_000.0, dt_s=0.1)])
    db.close_session(session)
    assert "PT1000 [°C]" in to_csv(db, session, tmp_path / "orphan.csv").read_text(encoding="utf-8")


def test_csv_export_honours_a_time_range(db, populated, tmp_path):
    options = ExportOptions(t_from=1_800_000_005.0, t_to=1_800_000_010.0)
    path = to_csv(db, populated, tmp_path / "slice.csv", options)
    rows = [line for line in path.read_text(encoding="utf-8").splitlines()
            if line and not line.startswith("#")]
    assert 40 < len(rows) < 60  # header plus ~51 samples


def test_excel_export_has_data_metadata_and_charts(db, populated, tmp_path):
    import openpyxl

    path = to_excel(db, populated, tmp_path / "out.xlsx")
    book = openpyxl.load_workbook(path)
    assert set(book.sheetnames) == {"Data", "Metadata", "Charts"}
    assert book["Data"].max_row == 201
    metadata = {row[0].value for row in book["Metadata"].iter_rows()}
    assert "Device serial" in metadata


def test_json_export_is_self_describing(db, populated, tmp_path):
    import json

    path = to_json(db, populated, tmp_path / "out.json")
    document = json.loads(path.read_text(encoding="utf-8"))
    assert document["format"] == "tjiptemp-export/1"
    assert {c["key"] for c in document["channels"]} == {"pt1000", "typek", "v_bat"}
    assert document["data"]["typek"][50] is None, "NaN must export as null, not 0"
    assert len(document["time_utc"]) == 200


def test_summary_table_counts_missing_samples(db, populated):
    table = summary_table(db, populated)
    typek = table[table["Channel"] == "Type K"].iloc[0]
    assert typek["Missing"] == 1
    assert typek["N"] == 199
    assert 22.0 < typek["Mean"] < 23.0


def test_resampling_does_not_fill_across_a_gap(db):
    """Resampling must leave a dropout empty rather than inventing values."""
    db.upsert_device("TJIP-G", {})
    session = db.create_session("TJIP-G", [int(Ch.PT1000)])
    left = np.full((10, 1), 20.0, dtype=np.float32)
    right = np.full((10, 1), 30.0, dtype=np.float32)
    db.append_blocks(session, [
        pack_block(0, left, t0_utc=0.0, dt_s=1.0),
        pack_block(100, right, t0_utc=100.0, dt_s=1.0),
    ])
    db.close_session(session)

    from tjiptemp.storage.export import load_frame

    frame, _, _ = load_frame(db, session, ExportOptions(resample_s=1.0))
    assert frame.isna().sum().sum() > 80, "the 90-second dropout must stay empty"


def test_figure_renders(db, populated, tmp_path):
    import matplotlib

    matplotlib.use("Agg")
    from tjiptemp.storage.export import to_figure_file

    path = to_figure_file(db, populated, tmp_path / "out.png")
    assert path.exists() and path.stat().st_size > 5000


def test_unknown_export_format_is_refused(db, populated, tmp_path):
    from tjiptemp.storage.export import ExportError, export

    with pytest.raises(ExportError, match="Supported"):
        export(db, populated, tmp_path / "out.xyz")
